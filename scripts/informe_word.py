"""Arma el informe con el formato del Taller 02: informe/informe.md → Word (.docx) y PDF.

    uv run python scripts/informe_word.py              # → informe/Informe_Taller03_Ballesteros_Alvarez_Ludena.docx y .pdf
    uv run python scripts/informe_word.py --revisar    # solo lista lo PENDIENTE y lo que falta incluir

El formato es el del informe entregado en el Taller 02: portada con el escudo de la USFQ, Calibri 11 justificado,
partes numeradas, la consigna del enunciado en cursiva con viñeta, leyendas «Ilustración N.» y «Tabla N.»,
recuadros de «Conclusión:» y número de página abajo a la derecha. El PDF se exporta con LibreOffice si está
instalado (`soffice`); si no, se abre el .docx en Word y se exporta como PDF.

Convenciones de informe.md:
    # 1. Título                      parte numerada (### y #### dentro de los archivos incluidos)
    ## Título / ### Título           subtítulos
    > *Consigna …*                   la consigna del enunciado, en cursiva, con viñeta
    > **Conclusión:** …              recuadro de conclusión
    *Ilustración: …*  /  *Tabla: …*  justo después de un bloque de código, una imagen o una tabla: su leyenda
    ![leyenda](ruta)                 figura (ruta relativa a la raíz del repositorio), con su leyenda
    incluir / incluir-codigo         comentarios HTML que insertan un archivo generado (ver scripts/informe_pdf.py):
                                     las tablas no se transcriben, se incluyen desde lo que escriben los scripts
    [PENDIENTE: …]                   lo que falta escribir: se resalta en naranja y se lista al generar
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import shutil
import subprocess
import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "scripts"))

import informe_pdf  # noqa: E402
from informe_pdf import PENDIENTE  # noqa: E402


def armar(md: str):
    """Como informe_pdf.armar, pero sin el título de primer nivel de cada Markdown incluido."""
    original = informe_pdf.bajar_titulos

    def sin_titulo(texto: str, niveles: int = 2) -> str:
        return original(re.sub(r"\A\s*# [^\n]*\n", "", texto), niveles)
    informe_pdf.bajar_titulos = sin_titulo
    try:
        return informe_pdf.armar(md)
    finally:
        informe_pdf.bajar_titulos = original

TITULO = "Taller 3 v2—Un solver multiagente con GraphRAG: del enunciado en PDF al reporte medido."
AUTORES = "Jessica Ballesteros, Miguel Álvarez, Darlyn Ludeña"
REPO = "Repositorio: github.com/akroasis0301/taller-03-opcional-alvarez-ballesteros-lude-a"
NOMBRE = "Informe_Taller03_Ballesteros_Alvarez_Ludena"
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
         "noviembre", "diciembre"]

AZUL_TITULO = "0F4761"      # Heading 2/3 del Taller 02
FONDO_CONCLUSION = "DCEAF7"
BORDE_CONCLUSION = "0A2F40"
FONDO_CABECERA = "DAE9F7"
FONDO_CODIGO = "F2F2F2"
ANCHO_TEXTO_CM = 15.9        # A4 con márgenes de 2,54 cm


# ============================================================================ Word de bajo nivel
def _docx():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    return Document, WD_ALIGN_PARAGRAPH, OxmlElement, qn, Cm, Pt, RGBColor


class Informe:
    def __init__(self):
        Document, self.AL, self.Ox, self.qn, self.Cm, self.Pt, self.RGB = _docx()
        self.doc = Document()
        self.n_ilustracion = 0
        self.n_tabla = 0
        s = self.doc.sections[0]
        s.page_width, s.page_height = self.Cm(21.0), self.Cm(29.7)
        for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, lado, self.Cm(2.54))
        normal = self.doc.styles["Normal"]
        normal.font.name, normal.font.size = "Calibri", self.Pt(11)
        normal.element.rPr.rFonts.set(self.qn("w:eastAsia"), "Calibri")
        normal.paragraph_format.space_after = self.Pt(6)
        normal.paragraph_format.line_spacing = 1.08
        self._pie()

    # ---------------------------------------------------------------- texto con formato en línea
    def runs(self, p, texto: str, base: dict | None = None, tam: float | None = None):
        """**negrita**, *cursiva*, `código` y [PENDIENTE: …] dentro de un párrafo."""
        base = base or {}
        patron = re.compile(r"(\[PENDIENTE:[^\]]*\]|\*\*[^*]+\*\*|`[^`]+`|(?<![\w*])\*[^*\s][^*]*\*(?![\w*]))")
        for trozo in patron.split(texto):
            if not trozo:
                continue
            fmt = dict(base)
            if trozo.startswith("[PENDIENTE:"):
                fmt.update(bold=True, pendiente=True)
            elif trozo.startswith("**") and trozo.endswith("**") and len(trozo) > 4:
                trozo, fmt["bold"] = trozo[2:-2], True
            elif trozo.startswith("`") and trozo.endswith("`") and len(trozo) > 2:
                trozo, fmt["code"] = trozo[1:-1], True
            elif trozo.startswith("*") and trozo.endswith("*") and len(trozo) > 2:
                trozo, fmt["italic"] = trozo[1:-1], not base.get("italic")
            # una negrita o un código dentro de una negrita o cursiva se parte otra vez
            if fmt != base and not fmt.get("code") and not fmt.get("pendiente") and \
                    (trozo.count("**") >= 2 or trozo.count("`") >= 2):
                self.runs(p, trozo, fmt, tam)
                continue
            r = p.add_run(trozo)
            r.bold = fmt.get("bold") or None
            r.italic = fmt.get("italic") or None
            if fmt.get("code"):
                r.font.name = "Consolas"
                r._element.rPr.rFonts.set(self.qn("w:eastAsia"), "Consolas")
                r.font.size = self.Pt((tam or 11) - 1.5)
            elif tam:
                r.font.size = self.Pt(tam)
            if fmt.get("pendiente"):
                r.font.color.rgb = self.RGB(0x9A, 0x34, 0x12)
                self._sombra_run(r, "FFEDD5")
        return p

    def _sombra_run(self, r, color):
        shd = self.Ox("w:shd")
        shd.set(self.qn("w:val"), "clear"), shd.set(self.qn("w:color"), "auto"), shd.set(self.qn("w:fill"), color)
        r._element.get_or_add_rPr().append(shd)

    def _sombra_celda(self, celda, color):
        tcPr = celda._tc.get_or_add_tcPr()
        shd = self.Ox("w:shd")
        shd.set(self.qn("w:val"), "clear"), shd.set(self.qn("w:color"), "auto"), shd.set(self.qn("w:fill"), color)
        tcPr.append(shd)

    def _bordes(self, tabla, color="000000", tam=4):
        tblPr = tabla._tbl.tblPr
        bordes = self.Ox("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = self.Ox(f"w:{lado}")
            b.set(self.qn("w:val"), "single"), b.set(self.qn("w:sz"), str(tam))
            b.set(self.qn("w:space"), "0"), b.set(self.qn("w:color"), color)
            bordes.append(b)
        tblPr.append(bordes)

    def _ancho_completo(self, t):
        layout = self.Ox("w:tblLayout")
        layout.set(self.qn("w:type"), "fixed")
        t._tbl.tblPr.append(layout)
        t.autofit = False
        t.columns[0].width = self.Cm(ANCHO_TEXTO_CM)
        c = t.rows[0].cells[0]
        c.width = self.Cm(ANCHO_TEXTO_CM)
        return c

    def _pie(self):
        p = self.doc.sections[0].footer.paragraphs[0]
        p.alignment = self.AL.RIGHT
        r = p.add_run()
        for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
            if tipo:
                f = self.Ox("w:fldChar")
                f.set(self.qn("w:fldCharType"), tipo)
                r._element.append(f)
            else:
                t = self.Ox("w:instrText")
                t.set(self.qn("xml:space"), "preserve")
                t.text = texto
                r._element.append(t)

    # ---------------------------------------------------------------- bloques
    def portada(self, fecha: str):
        logo = RAIZ / "informe" / "usfq_logo.png"
        if logo.exists():
            p = self.doc.add_paragraph()
            p.alignment = self.AL.CENTER
            p.add_run().add_picture(str(logo), width=self.Cm(4.05))
        for texto, tam, negrita in [("UNIVERSIDAD SAN FRANCISCO DE QUITO", 14, True),
                                    ("Maestría en Inteligencia Artificial", 14, True),
                                    ("IA Generativa y Agentes", 14, True), (TITULO, 13, False), (AUTORES, 13, False),
                                    (fecha, 13, False), (REPO, 13, False)]:
            p = self.doc.add_paragraph()
            p.alignment = self.AL.CENTER
            p.paragraph_format.space_after = self.Pt(10)
            r = p.add_run(texto)
            r.bold, r.font.size = negrita or None, self.Pt(tam)

    def titulo(self, texto: str, nivel: int):
        p = self.doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        if nivel == 1:                                  # «1. Parte 0 — …», como las partes del Taller 02
            p.paragraph_format.space_before = self.Pt(14)
            p.paragraph_format.left_indent = self.Cm(0.63)
            self.runs(p, texto, {"bold": True}, 12)
            return
        tam = {2: 14, 3: 12.5}.get(nivel, 11.5)
        p.paragraph_format.space_before = self.Pt(10 if nivel <= 3 else 6)
        self.runs(p, texto, {"bold": nivel >= 4}, tam)
        for r in p.runs:
            if r.font.color.rgb is None:
                r.font.color.rgb = self.RGB.from_string(AZUL_TITULO)

    def parrafo(self, texto: str, sangria: float = 0, vineta: str | None = None, base: dict | None = None):
        p = self.doc.add_paragraph()
        p.alignment = self.AL.JUSTIFY
        if vineta:
            p.paragraph_format.left_indent = self.Cm(0.63 + sangria)
            p.paragraph_format.first_line_indent = self.Cm(-0.5)
            p.paragraph_format.space_after = self.Pt(3)
            p.add_run(f"{vineta}\t")
            p.paragraph_format.tab_stops.add_tab_stop(self.Cm(0.63 + sangria))
        return self.runs(p, texto, base)

    def consigna(self, texto: str):
        """La consigna del enunciado: viñeta, cursiva, con el título de la sección en negrita."""
        texto = texto.strip().strip("*").strip()
        m = re.match(r"^(.{3,90}?\.)\s+(.*)$", texto, re.S)
        cabeza, resto = (m.group(1), m.group(2)) if m else ("", texto)
        p = self.parrafo("", vineta="•")
        p.paragraph_format.space_after = self.Pt(8)
        if cabeza:
            self.runs(p, cabeza + " ", {"bold": True, "italic": True})
        self.runs(p, resto, {"italic": True})

    def recuadro(self, texto: str, fondo: str = FONDO_CONCLUSION, borde: str = BORDE_CONCLUSION):
        t = self.doc.add_table(rows=1, cols=1)
        self._bordes(t, borde, 8)
        c = self._ancho_completo(t)
        self._sombra_celda(c, fondo)
        p = c.paragraphs[0]
        p.alignment = self.AL.JUSTIFY
        self.runs(p, texto)
        self.doc.add_paragraph().paragraph_format.space_after = self.Pt(2)

    def codigo(self, texto: str):
        t = self.doc.add_table(rows=1, cols=1)
        self._bordes(t, "BFBFBF", 4)
        c = self._ancho_completo(t)
        self._sombra_celda(c, FONDO_CODIGO)
        lineas = texto.rstrip("\n").splitlines() or [""]
        p = c.paragraphs[0]
        p.paragraph_format.space_after = self.Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for i, linea in enumerate(lineas):
            r = p.add_run(linea)
            r.font.name, r.font.size = "Consolas", self.Pt(7.5)
            r._element.rPr.rFonts.set(self.qn("w:eastAsia"), "Consolas")
            if i < len(lineas) - 1:
                r.add_break()

    def leyenda(self, tipo: str, texto: str):
        if tipo == "Tabla":
            self.n_tabla += 1
            n = self.n_tabla
        else:
            self.n_ilustracion += 1
            n = self.n_ilustracion
        p = self.doc.add_paragraph()
        p.alignment = self.AL.CENTER
        p.paragraph_format.space_before = self.Pt(3)
        self.runs(p, f"{tipo} {n}. {texto}", {"italic": True}, 10)

    def imagen(self, ruta: Path):
        from PIL import Image
        ancho, alto = Image.open(ruta).size
        ancho_cm = min(ANCHO_TEXTO_CM, 16.0)
        if alto / ancho > 0.8:                          # figuras altas: que no ocupen la página entera
            ancho_cm = min(ancho_cm, 13.0 * ancho / alto)
        p = self.doc.add_paragraph()
        p.alignment = self.AL.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(ruta), width=self.Cm(ancho_cm))

    def tabla(self, filas: list[list[str]]):
        ncol = max(len(f) for f in filas)
        t = self.doc.add_table(rows=len(filas), cols=ncol)
        self._bordes(t)
        tam = 9 if ncol <= 4 else 8 if ncol <= 7 else 7
        # anchos fijos, proporcionales al texto más largo de cada columna (sin esto LibreOffice las reparte mal)
        # mínimo: la palabra más larga de la columna sin partir; deseado: el texto más largo (hasta 40 caracteres)
        cm_por_caracter, relleno = 0.0185 * tam, 0.3
        minimo, deseado = [], []
        for j in range(ncol):
            textos = [re.sub(r"[*`]", "", f[j]) if j < len(f) else "" for f in filas]
            palabra = max((len(w) for t_ in textos for w in t_.split()), default=1)
            minimo.append(min(palabra, 12) * cm_por_caracter + relleno)
            deseado.append(max(min(max(len(t_) for t_ in textos), 40) * cm_por_caracter + relleno, minimo[-1]))
        if sum(deseado) <= ANCHO_TEXTO_CM:
            anchos = [d * ANCHO_TEXTO_CM / sum(deseado) for d in deseado]
        elif sum(minimo) >= ANCHO_TEXTO_CM:
            anchos = [m * ANCHO_TEXTO_CM / sum(minimo) for m in minimo]
        else:
            sobra, extra = ANCHO_TEXTO_CM - sum(minimo), [d - m for d, m in zip(deseado, minimo)]
            anchos = [m + sobra * e / sum(extra) for m, e in zip(minimo, extra)]
        tblPr = t._tbl.tblPr
        layout = self.Ox("w:tblLayout")
        layout.set(self.qn("w:type"), "fixed")
        tblPr.append(layout)
        t.autofit = False
        for j, col in enumerate(t.columns):
            col.width = self.Cm(anchos[j])
        for i, fila in enumerate(filas):
            for j in range(ncol):
                c = t.rows[i].cells[j]
                c.width = self.Cm(anchos[j])
                p = c.paragraphs[0]
                p.paragraph_format.space_after = self.Pt(0)
                p.paragraph_format.line_spacing = 1.0
                texto = fila[j] if j < len(fila) else ""
                self.runs(p, texto, {"bold": i == 0}, tam)
                if i == 0:
                    self._sombra_celda(c, FONDO_CABECERA)
                    p.alignment = self.AL.CENTER
                if len(filas) <= 12 and i < len(filas) - 1:     # una tabla corta no se parte entre páginas
                    p.paragraph_format.keep_with_next = True
            trPr = t.rows[i]._tr.get_or_add_trPr()
            if i == 0:                                          # la cabecera se repite si la tabla sigue
                trPr.append(self.Ox("w:tblHeader"))
            trPr.append(self.Ox("w:cantSplit"))
        self.doc.add_paragraph().paragraph_format.space_after = self.Pt(0)


# ============================================================================ Markdown → bloques
LEYENDA = re.compile(r"^\*(Ilustración|Tabla):\s*(.+)\*\s*$")
IMAGEN = re.compile(r"^!\[([^\]]*)\]\(([^)\s]+)\)\s*$")


def celdas(linea: str) -> list[str]:
    return [c.strip() for c in re.split(r"(?<!\\)\|", linea.strip().strip("|"))]


def convertir(md: str, inf: Informe) -> list[str]:
    md = re.sub(r"<!--.*?-->", "", md, flags=re.S)
    lineas = md.splitlines()
    faltan_img: list[str] = []
    i = 0

    def leyenda_siguiente(j: int) -> tuple[str, str] | None:
        while j < len(lineas) and not lineas[j].strip():
            j += 1
        if j < len(lineas):
            m = LEYENDA.match(lineas[j].strip())
            if m:
                lineas[j] = ""
                return m.group(1), m.group(2)
        return None

    parrafo: list[str] = []

    def vaciar():
        if parrafo:
            texto = " ".join(x.strip() for x in parrafo)
            m_ = re.fullmatch(r"_(.+)_", texto)
            inf.parrafo(m_.group(1) if m_ else texto, base={"italic": True} if m_ else None)
            parrafo.clear()

    while i < len(lineas):
        linea = lineas[i]
        s = linea.strip()
        if not s or s == "---":
            vaciar()
            i += 1
            continue
        if s.startswith("```"):
            vaciar()
            j = i + 1
            while j < len(lineas) and not lineas[j].strip().startswith("```"):
                j += 1
            inf.codigo("\n".join(lineas[i + 1:j]))
            ley = leyenda_siguiente(j + 1)
            if ley:
                inf.leyenda(*ley)
            else:
                inf.doc.add_paragraph().paragraph_format.space_after = 0
            i = j + 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            vaciar()
            inf.titulo(m.group(2), len(m.group(1)))
            i += 1
            continue
        m = IMAGEN.match(s)
        if m:
            vaciar()
            ruta = RAIZ / m.group(2)
            if ruta.exists():
                inf.imagen(ruta)
                ley = leyenda_siguiente(i + 1)
                inf.leyenda(*(ley or ("Ilustración", m.group(1))))
            else:
                faltan_img.append(m.group(2))
                inf.parrafo(f"[PENDIENTE: falta la imagen {m.group(2)}]")
            i += 1
            continue
        if s.startswith("|"):
            vaciar()
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                c = celdas(lineas[i])
                if not all(re.fullmatch(r":?-{2,}:?", x) for x in c if x):
                    filas.append(c)
                i += 1
            inf.tabla(filas)
            ley = leyenda_siguiente(i)
            if ley:
                inf.doc.paragraphs[-1]._element.getparent().remove(inf.doc.paragraphs[-1]._element)
                inf.leyenda(*ley)
            continue
        if s.startswith(">"):
            vaciar()
            bloque = []
            while i < len(lineas) and lineas[i].strip().startswith(">"):
                bloque.append(lineas[i].strip()[1:].strip())
                i += 1
            texto = " ".join(bloque)
            if texto.startswith("**Conclusión:**") or texto.startswith("**Conclusión**"):
                inf.recuadro(texto)
            else:
                inf.consigna(texto)
            continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", linea)
        if m and not LEYENDA.match(s):
            vaciar()
            sangria = 0.63 * (len(m.group(1)) // 2)
            item = [m.group(3)]
            i += 1
            while i < len(lineas) and lineas[i].strip() and lineas[i].startswith(" ") and \
                    not re.match(r"^\s*([-*]|\d+\.)\s+", lineas[i]):
                item.append(lineas[i].strip())
                i += 1
            vineta = "•" if m.group(2) in "-*" else m.group(2)
            inf.parrafo(" ".join(item), sangria, vineta)
            continue
        m = LEYENDA.match(s)
        if m:
            vaciar()
            inf.leyenda(m.group(1), m.group(2))
            i += 1
            continue
        parrafo.append(linea)
        i += 1
    vaciar()
    return faltan_img


def fecha_es(d: dt.date) -> str:
    return f"{d.day} de {MESES[d.month - 1]} de {d.year}"


def a_pdf(docx: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if not soffice and mac.exists():
        soffice = str(mac)
    if not soffice:
        return None
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                   check=True, capture_output=True, timeout=300)
    pdf = docx.with_suffix(".pdf")
    return pdf if pdf.exists() else None


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--entrada", default=str(RAIZ / "informe" / "informe.md"))
    ap.add_argument("--salida", default=str(RAIZ / "informe" / f"{NOMBRE}.docx"))
    ap.add_argument("--fecha", help="fecha de la portada, AAAA-MM-DD (por omisión, hoy)")
    ap.add_argument("--revisar", action="store_true")
    a = ap.parse_args()

    md, faltan = armar(Path(a.entrada).read_text(encoding="utf-8"))
    pendientes = PENDIENTE.findall(re.sub(r"<!--.*?-->", "", md, flags=re.S))
    if a.revisar:
        for etiqueta, lista in [("archivos por incluir que no existen", faltan), ("PENDIENTES", pendientes)]:
            print(f"{etiqueta}: {len(lista)}")
            for x in lista:
                print(f"   - {x[:150]}")
        return 0
    inf = Informe()
    fecha = dt.date.fromisoformat(a.fecha) if a.fecha else dt.date.today()
    inf.portada(fecha_es(fecha))
    faltan_img = convertir(md, inf)
    salida = Path(a.salida)
    inf.doc.core_properties.title = TITULO
    inf.doc.core_properties.author = AUTORES
    inf.doc.save(salida)
    print(f"→ {salida}  ({inf.n_ilustracion} ilustraciones, {inf.n_tabla} tablas)")
    pdf = a_pdf(salida)
    print(f"→ {pdf}" if pdf else "PDF: no hay LibreOffice; abre el .docx en Word y exporta como PDF.")
    for etiqueta, lista in [("archivos por incluir que no existen", faltan), ("imágenes que no existen", faltan_img),
                            ("PENDIENTES", pendientes)]:
        print(f"{etiqueta}: {len(lista)}")
        for x in lista:
            print(f"   - {x[:150]}")
    if pendientes or faltan or faltan_img:
        print("AVISO: el informe tiene partes pendientes; no es la versión de entrega.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
